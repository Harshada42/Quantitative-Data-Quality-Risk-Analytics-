import mysql.connector
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

# --- 1. Connect to MySQL Database ---
print("Connecting to MySQL Database for reporting...")
db = mysql.connector.connect(
    host="127.0.0.1",
    user="root",
    password="",
    database="risk_analytics_db"
)

# --- 2. Query Executive Risk Data ---
query = """
    SELECT 
        transaction_id,
        customer_id,
        vendor_id,
        transaction_date,
        amount,
        payment_method,
        is_flagged_anomaly,
        anomaly_score
    FROM fact_transactions
    ORDER BY amount DESC;
"""

df = pd.read_sql(query, db)
db.close()
print(f"Fetched {len(df)} records from MySQL for report generation.")

# --- 3. Build Professional Excel Workbook ---
wb = Workbook()
ws = wb.active
ws.title = "Risk Analytics Summary"

# Ensure grid lines are visible
ws.views.sheetView[0].showGridLines = True

# Styling definitions
font_title = Font(name="Calibri", size=16, bold=True, color="1F4E78")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
font_data = Font(name="Calibri", size=11)
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
border_thin = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

# Soft Red styling for High-Risk Anomalies
fill_anomaly = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
font_anomaly = Font(name="Calibri", size=11, color="9C0006", bold=True)

# Add Title Block
ws.append([]) # Blank row
ws.append(["Executive Risk & Transaction Quality Report"])
ws.cell(row=2, column=1).font = font_title
ws.append([f"Generated automatically from MySQL via Python pipeline"])
ws.cell(row=3, column=1).font = Font(name="Calibri", size=9, italic=True, color="595959")
ws.append([]) # Blank row

# Add Table Headers
headers = [
    "Transaction ID", "Customer ID", "Vendor ID", 
    "Transaction Date", "Amount ($)", "Payment Method", 
    "Flagged Risk", "Anomaly Score"
]
ws.append(headers)

header_row_idx = 5
for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=header_row_idx, column=col_idx)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center

# Append Data Rows
for row_data in dataframe_to_rows(df, index=False, header=False):
    ws.append(row_data)

# Apply Formatting to Data Rows
start_row = 6
end_row = start_row + len(df) - 1

for row_idx in range(start_row, end_row + 1):
    is_anomaly = ws.cell(row=row_idx, column=7).value # is_flagged_anomaly column
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.font = font_data
        cell.border = border_thin
        
        # Format Currency column
        if col_idx == 5:
            cell.number_format = '$#,##0.00'
            cell.alignment = align_right
        elif col_idx in [1, 2, 3, 4, 6, 7, 8]:
            cell.alignment = align_center

        # Highlight High-Risk Anomalies
        if is_anomaly == 1:
            cell.fill = fill_anomaly
            if col_idx == 7:
                cell.font = font_anomaly

# Auto-adjust column widths
for col in ws.columns:
    max_length = max(len(str(cell.value or '')) for cell in col)
    col_letter = col[0].column_letter
    ws.column_dimensions[col_letter].width = max(max_length + 4, 12)

# Save Workbook
output_path = "data/Executive_Risk_Report.xlsx"
wb.save(output_path)
print(f"Professional Excel report successfully generated and saved to {output_path}!")